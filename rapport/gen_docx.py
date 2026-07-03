# -*- coding: utf-8 -*-
"""Build a properly-formatted Arabic Word (.docx) thesis from the student's
original .docx, applying the requested corrections + Moroccan psychology
academic formatting standards."""
import zipfile, re, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX_IN = "/root/.claude/uploads/d37b352c-c516-52d1-98da-8d676b1ba3f6/50dfb9dc-_____.docx"
OUT     = "/home/user/MP3Translator/rapport/rapport.docx"
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
ARFONT = 'Traditional Arabic'

# ================= parse original docx =================
def para_text(p): return ''.join(t.text or '' for t in p.iter(W+'t')).strip()
def para_style(p):
    ppr=p.find(W+'pPr')
    if ppr is None: return ''
    ps=ppr.find(W+'pStyle'); return ps.get(W+'val','') if ps is not None else ''
def dedup(t):
    n=len(t)
    return t[:n//2] if (n>8 and n%2==0 and t[:n//2]==t[n//2:]) else t

root=ET.fromstring(zipfile.ZipFile(DOCX_IN).read('word/document.xml'))
blocks=[]
for el in root.find(W+'body'):
    tag=el.tag.replace(W,'')
    if tag=='p':
        txt=dedup(para_text(el)); st=para_style(el)
        if txt: blocks.append(('P',st,txt))
    elif tag=='tbl':
        rows=[]
        for tr in el.findall(W+'tr'):
            rows.append([dedup(' '.join(para_text(p) for p in tc.findall(W+'p')).strip())
                         for tc in tr.findall(W+'tc')])
        blocks.append(('T',None,rows))

# ================= clean text helpers =================
def clean_cell(s):
    for ch in ['🔴','🟡','🟢','🔵']: s=s.replace(ch,'')
    return s.strip()
def clean_tex(t):
    t=re.sub(r'\\textbf\{(.+?)\}', r'⟦b⟧\1⟦/b⟧', t)
    t=re.sub(r'\\emph\{(.+?)\}', r'⟦b⟧\1⟦/b⟧', t)
    t=re.sub(r'\\(?:textenglish|en)\{(.+?)\}', r'\1', t)
    t=t.replace('\\noindent','').replace('\\par','').replace('\\,',' ').replace('\\ ',' ')
    t=t.replace('\\%','%').replace('\\&','&').replace('\\_','_').replace('\\#','#')
    t=t.replace('--','–')
    t=re.sub(r'\\[a-zA-Z]+\*?','',t)
    t=t.replace('{','').replace('}','')
    return re.sub(r'\s+',' ',t).strip()

def parse_tex(path):
    out=[]; buf=[]
    def flush():
        if buf:
            txt=clean_tex(' '.join(buf))
            if txt: out.append(('p',txt))
        buf.clear()
    for ln in open(path,encoding='utf-8'):
        s=ln.strip()
        if not s or s in ('\\medskip','\\par','\\bigskip'): flush(); continue
        if s.startswith('\\addcontentsline'): continue
        m=re.match(r'\\chapter\*?\{(.+)\}$',s)
        if m: flush(); out.append(('h1',clean_tex(m.group(1)))); continue
        m=re.match(r'\\section\*?\{(.+)\}$',s)
        if m: flush(); out.append(('h2',clean_tex(m.group(1)))); continue
        m=re.match(r'\\subsection\*?\{(.+)\}$',s)
        if m: flush(); out.append(('h3',clean_tex(m.group(1)))); continue
        if re.match(r'\\begin\{(enumerate|itemize)\}',s) or re.match(r'\\end\{(enumerate|itemize)\}',s):
            flush(); continue
        m=re.match(r'\\item\s*(.*)$',s)
        if m: flush(); out.append(('li',clean_tex(m.group(1)))); continue
        buf.append(s)
    flush()
    return out

NEW_INTRO=parse_tex('/home/user/MP3Translator/rapport/_intro.tex')
NEW_CH4  =parse_tex('/home/user/MP3Translator/rapport/_ch4.tex')

# ================= docx low-level helpers =================
def _set(el, tag, **attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items(): e.set(qn(k),v)
    el.append(e); return e

def style_run(run, size, bold=False):
    run.font.name=ARFONT; run.font.size=Pt(size); run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    rf.set(qn('w:cs'),ARFONT); rf.set(qn('w:ascii'),ARFONT); rf.set(qn('w:hAnsi'),ARFONT)
    _set(rpr,'w:szCs',**{'w:val':str(int(size*2))})
    _set(rpr,'w:rtl',**{'w:val':'1'})

def rtl_par(p):
    ppr=p._p.get_or_add_pPr(); _set(ppr,'w:bidi',**{'w:val':'1'})

def add_runs(p, text, size, bold_all=False):
    # split on ⟦b⟧ markers into bold/normal runs
    parts=re.split(r'(⟦b⟧|⟦/b⟧)', text)
    bold=bold_all
    for seg in parts:
        if seg=='⟦b⟧': bold=True; continue
        if seg=='⟦/b⟧': bold=bold_all; continue
        if seg=='': continue
        r=p.add_run(seg); style_run(r,size,bold)

def para(doc, text, size=14, bold=False, align='justify', indent=True, space_after=6):
    p=doc.add_paragraph(); rtl_par(p)
    pf=p.paragraph_format
    pf.line_spacing=1.5
    pf.space_after=Pt(space_after); pf.space_before=Pt(0)
    pf.alignment={'justify':WD_ALIGN_PARAGRAPH.JUSTIFY,'center':WD_ALIGN_PARAGRAPH.CENTER,
                  'right':WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if indent and align=='justify': pf.first_line_indent=Cm(1.27)
    add_runs(p, text, size, bold_all=bold)
    return p

def heading(doc, text, level):
    size={1:16,2:15,3:14}[level]
    p=doc.add_paragraph(); rtl_par(p)
    pf=p.paragraph_format; pf.line_spacing=1.0
    pf.space_before=Pt(12 if level==1 else 8); pf.space_after=Pt(6)
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    # keep heading in TOC via outline level
    ppr=p._p.get_or_add_pPr(); _set(ppr,'w:outlineLvl',**{'w:val':str(level-1)})
    add_runs(p, text, size, bold_all=True)
    if level==1:
        pf.space_before=Pt(18)
    return p

def add_table(doc, rows, caption=None, tnum=None):
    rows=[r for r in rows if any(c.strip() for c in r)]
    if not rows: return
    ncols=max(len(r) for r in rows)
    rows=[r+['']*(ncols-len(r)) for r in rows]
    if caption:
        cap='جدول' + (f' ({tnum}): ' if tnum else ': ') + caption
        p=para(doc, cap, size=12, bold=True, align='center', indent=False, space_after=2)
    t=doc.add_table(rows=len(rows), cols=ncols)
    t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.allow_autofit=True
    # RTL table + fit to 100% width
    tblPr=t._tbl.tblPr
    _set(tblPr,'w:bidiVisual')
    tw=tblPr.find(qn('w:tblW'))
    if tw is None: tw=OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:type'),'pct'); tw.set(qn('w:w'),'5000')
    fsize=9 if ncols>=6 else (10 if ncols>=4 else 12)
    for ri,row in enumerate(rows):
        for ci in range(ncols):
            cell=t.cell(ri,ci); cell.text=''
            cp=cell.paragraphs[0]; rtl_par(cp)
            cp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.line_spacing=1.0; cp.paragraph_format.space_after=Pt(0)
            add_runs(cp, clean_cell(row[ci]).replace('✓','✔'), fsize, bold_all=(ri==0))
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def field(par, instr):
    r=par.add_run(); fc=OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),'begin'); r._r.append(fc)
    r2=par.add_run(); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=instr; r2._r.append(it)
    r3=par.add_run(); fc2=OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'),'separate'); r3._r.append(fc2)
    r4=par.add_run('...'); style_run(r4,12)
    r5=par.add_run(); fc3=OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'),'end'); r5._r.append(fc3)

# ================= build document =================
doc=Document()
# --- page setup + default style
sec=doc.sections[0]
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5)
sec.left_margin=Cm(3); sec.right_margin=Cm(3)
# RTL section
sec._sectPr.append(OxmlElement('w:bidi'))
normal=doc.styles['Normal']
normal.font.name=ARFONT; normal.font.size=Pt(14)
rpr=normal.element.get_or_add_rPr()
rf=OxmlElement('w:rFonts'); rf.set(qn('w:cs'),ARFONT); rf.set(qn('w:ascii'),ARFONT); rf.set(qn('w:hAnsi'),ARFONT); rpr.append(rf)
# footer page number (centered)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; field(fp,' PAGE ')
# update fields (TOC) on open
try:
    st=doc.settings.element; _set(st,'w:updateFields',**{'w:val':'true'})
except Exception: pass

# --- COVER (kept as the student's original, city = Tétouan) ---
cover=['شعبة علم النفس الإكلينيكي','جامعة عبد المالك السعدي','كلية الآداب والعلوم الإنسانية بتطوان','',
       'بحث لنيل شهادة الإجازة في علم النفس الإكلينيكي','','بعنوان:','',
       'إدراك أسلوب المعاملة الوالدية المتسلط وعلاقته بتقدير الذات لدى المراهقين (15–17 سنة)','دراسة كيفية','','',
       'إعداد الطالبة: لطيفة أديب','رقم التسجيل: 23068742','','تحت إشراف الأستاذ: د. بدر الدين الزيدي','','',
       'المستوى: السنة الثالثة من الإجازة – علم النفس الإكلينيكي','السنة الجامعية: 2025 – 2026']
for i,l in enumerate(cover):
    big = (l.startswith('إدراك'))
    para(doc, l, size=(18 if big else 14), bold=(i==0 or big or l=='بعنوان:'),
         align='center', indent=False, space_after=(4 if l else 10))
doc.add_page_break()

# --- FRONT MATTER: TOC + LoT fields ---
heading(doc,'فهرس المحتويات',1)
tp=doc.add_paragraph(); rtl_par(tp); tp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field(tp,'TOC \\o "1-3" \\h \\z \\u')
doc.add_page_break()
heading(doc,'لائحة الجداول',1)
lp=doc.add_paragraph(); rtl_par(lp); lp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field(lp,'TOC \\h \\z \\c "جدول"')
doc.add_page_break()

# ================= main conversion loop =================
SEC_PREF=('أولاً','أولًا','ثانياً','ثانيًا','ثالثاً','ثالثًا','رابعاً','رابعًا','خامساً','خامسًا',
          'سادساً','سادسًا','سابعاً','سابعًا','ثامناً','ثامنًا','تاسعاً','تاسعًا','عاشراً')
BED=('البعد الأول','البعد الثاني','البعد الثالث','البعد الرابع','البعد الخامس')
MEHWAR=('المحور الأول','المحور الثاني','المحور الثالث')
SUBQ=('السؤال الأول','السؤال الثاني','السؤال الافتتاحي','السؤال الختامي','سؤال افتتاحي','سؤال ختامي')
def clean_lead(t): return t.lstrip('.-–—0123456789 ').strip()
def hlevel(t):
    tt=t.strip()
    if len(tt)<70 and tt.startswith(SEC_PREF): return (2,tt)
    if len(tt)<70 and any(tt.startswith(b) for b in BED): return (2,tt)
    if len(tt)<85 and any(tt.startswith(m) for m in MEHWAR): return (2,tt)
    if tt.startswith('الملحق'): return (2,tt)
    if len(tt)<70 and any(tt.startswith(q) for q in SUBQ): return (3,tt)
    if tt in ('تمهيد','تمهيد الفصل'): return (2,'تمهيد')
    if tt in ('الأسئلة الفرعية','السؤال المركزي'): return (3,tt)
    if 'التعريف الإجرائي' in tt and len(tt)<45: return (3,tt)
    if tt.startswith('مفهوم الإدراك'): return (3,tt)
    if tt.startswith('. ') and len(tt)<80: return (3,clean_lead(tt))
    if re.match(r'^\d+[.\-]\s',tt) and len(tt)<45: return (3,clean_lead(tt))
    return (None,None)
RUNIN=('الجواب:','السؤال:','المؤشر المستخلص:','المؤشر المستهدف:','البعد المرتبط:','السند النظري:',
       'الأسئلة:','وظيفته:','الملاحظة التحليلية','الاتجاه العام','الحالات الخاصة','القراءة التحليلية',
       'خلاصة الفصل','ملاحظة منهجية','ملاحظة:','الكلمات المفتاحية:')

def emit_blocks(bl):
    for k,t in bl:
        if k=='h1': heading(doc,t,1)
        elif k=='h2': heading(doc,t,2)
        elif k=='h3': heading(doc,t,3)
        elif k=='li':
            p=para(doc,'•  '+t,size=14,align='justify',indent=False,space_after=3)
        else: para(doc,t,size=14)

i=0; n=len(blocks); pending=None; mode='skip_intro'; emitted_ch2=False
in_appendix=False; refs_mode=False; pipeline_done=False
def is_cap(t): return t.strip().startswith('جدول (')
def strip_cap(t):
    t=t.strip(); return t.split(':',1)[1].strip() if ':' in t else t
tnum=0
while i<n:
    kind,st,val=blocks[i]
    if kind=='T':
        global_cap=pending; pending=None
        cap=global_cap
        tnum2=None
        if cap:
            tnum+=1; tnum2=tnum
        add_table(doc,val,cap,tnum2); i+=1; continue
    t=val.strip()
    if mode=='skip_intro':
        if t.startswith('الفصل الاول') or t.startswith('الفصل الأول'):
            mode='normal'; emit_blocks(NEW_INTRO); heading(doc,'الفصل الأول: الإطار النظري للدراسة',1)
        i+=1; continue
    if t.startswith('الفصل الرابع'):
        emit_blocks(NEW_CH4)
        i+=1
        while i<n and not (blocks[i][0]=='P' and blocks[i][2].strip().startswith('الخاتمة العامة')): i+=1
        continue
    if (not emitted_ch2) and (t.startswith('4. هدف البحث') or ('هدف البحث' in t and len(t)<25)):
        heading(doc,'الفصل الثاني: الإشكالية وإجراءات الدراسة المنهجية',1)
        heading(doc,'تمهيد',2)
        para(doc,'تتناول الدراسة العلاقة بين إدراك المراهق لأسلوب المعاملة الوالدية المتسلط وتقديره '
                 'لذاته، انطلاقًا من أنّ أثر الممارسات الوالدية لا يتحدّد بوجودها في ذاتها وإنما بالكيفية '
                 'التي يدركها بها الأبناء ويفسّرونها، وهو ما اقتضى المقاربة الكيفية.')
        heading(doc,'هدف البحث',2); emitted_ch2=True; i+=1; continue
    if t.startswith('ؤ '): i+=1; continue
    if ('الإطار النظري للدراسةالإطار' in val or 'إجراء المقابلات وجمع البياناتإجراء' in val
        or 'تحديد المحور الأول وأبعاده (أسلوب' in val):
        if not pipeline_done:
            for step in ['بناء الإطار النظري للدراسة','تحديد المحور الأول وأبعاده (أسلوب المعاملة الوالدية المتسلط)',
                'تحديد المحور الثاني وأبعاده (تقدير الذات)','بناء دليل المقابلة شبه الموجّهة',
                'اختيار العيّنة القصدية وإجراء الفرز القبلي','إجراء المقابلات وجمع البيانات',
                'تفريغ المقابلات والقراءة المتكرّرة','ترميز البيانات وتصنيف الترميزات',
                'تحليل المضمون (التكرار – الاتجاه العام – الحالات الخاصة)','عرض النتائج وتفسيرها']:
                para(doc,'•  '+step,size=14,indent=False,space_after=2)
            pipeline_done=True
        i+=1; continue
    if is_cap(t): pending=strip_cap(t); i+=1; continue
    if t.startswith('الملحق (') and 'جدول' in t:
        pending=strip_cap(t.split('جدول',1)[1]); heading(doc,t.split(':',1)[0].strip(),3); i+=1; continue
    if t.startswith('الفصل الثالث'): heading(doc,'الفصل الثالث: عرض النتائج وتحليلها',1); i+=1; continue
    if t.startswith('الخاتمة العامة'): heading(doc,'الخاتمة العامة',1); i+=1; continue
    if t.startswith('الملاحق') and len(t)<12:
        heading(doc,'الملاحق',1); in_appendix=True; i+=1; continue
    if t.startswith('قائمة المراجع'): heading(doc,'قائمة المراجع',1); refs_mode=True; i+=1; continue
    if t.startswith('لائحة الجداول'): i+=1; continue
    if re.match(r'^المشارك\s+م\d+',t) or re.match(r'^ملخص الحالة',t):
        heading(doc,t,3); i+=1; continue
    if t.startswith('الرمز:'):
        e=t
        for kk in ['الجنس:','العمر:','المستوى الدراسي:']: e=e.replace(kk,'   |   '+kk)
        para(doc,e,size=13,bold=True,indent=False); i+=1; continue
    lvl,htext=hlevel(t)
    if lvl:
        heading(doc,htext,3 if in_appendix else lvl); i+=1; continue
    if st=='Paragraphedeliste':
        para(doc,'•  '+t,size=14,indent=False,space_after=3); i+=1; continue
    # runin bold label
    if (t.endswith(':') and len(t)<45) or any(t.startswith(k) for k in RUNIN):
        if ':' in t and not t.endswith(':'):
            lab,rest=t.split(':',1); para(doc,f'⟦b⟧{lab.strip()}:⟦/b⟧ {rest.strip()}',size=14)
        else: para(doc,f'⟦b⟧{t}⟦/b⟧',size=14,indent=False)
        i+=1; continue
    if t.startswith('- '): t=t[2:].strip()
    # references get size 12
    para(doc,t,size=(12 if refs_mode else 14),
         indent=(not refs_mode), space_after=(4 if refs_mode else 6))
    i+=1

doc.save(OUT)
print('SAVED',OUT,'paragraphs:',len(doc.paragraphs),'tables:',len(doc.tables))
