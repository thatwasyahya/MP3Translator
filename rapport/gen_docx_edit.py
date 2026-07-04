# -*- coding: utf-8 -*-
"""Edit the student's ORIGINAL .docx in place: fix cover (Faculté des Sciences,
Tétouan), swap in the new introduction + rebuilt Chapter 4, make every table fit
the page (RTL), and normalise the layout to the Moroccan academic standards."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/root/.claude/uploads/d37b352c-c516-52d1-98da-8d676b1ba3f6/50dfb9dc-_____.docx"
OUT = "/home/user/MP3Translator/rapport/rapport.docx"
AR  = 'Traditional Arabic'

# ---------- text cleaning for the new intro / ch4 (.tex -> plain) ----------
def clean_tex(t):
    t=re.sub(r'\\textbf\{(.+?)\}', r'⟦b⟧\1⟦/b⟧', t)
    t=re.sub(r'\\emph\{(.+?)\}',  r'⟦b⟧\1⟦/b⟧', t)
    t=re.sub(r'\\(?:textenglish|en)\{(.+?)\}', r'\1', t)
    t=t.replace('\\noindent','').replace('\\par','').replace('\\,',' ').replace('\\ ',' ')
    t=t.replace('\\%','%').replace('\\&','&').replace('\\_','_').replace('\\#','#')
    t=t.replace('--','–')
    t=re.sub(r'\\[a-zA-Z]+\*?','',t).replace('{','').replace('}','')
    return re.sub(r'\s+',' ',t).strip()

def parse_tex(path):
    out=[]; buf=[]
    def flush():
        if buf:
            x=clean_tex(' '.join(buf));  out.append(('p',x)) if x else None
        buf.clear()
    for ln in open(path,encoding='utf-8'):
        s=ln.strip()
        if not s or s in ('\\medskip','\\par','\\bigskip'): flush(); continue
        if s.startswith('\\addcontentsline'): continue
        for cmd,tag in (('chapter','h1'),('section','h2'),('subsection','h3')):
            m=re.match(r'\\%s\*?\{(.+)\}$'%cmd,s)
            if m: flush(); out.append((tag,clean_tex(m.group(1)))); break
        else:
            if re.match(r'\\(begin|end)\{(enumerate|itemize)\}',s): flush(); continue
            m=re.match(r'\\item\s*(.*)$',s)
            if m: flush(); out.append(('li',clean_tex(m.group(1)))); continue
            buf.append(s)
    flush(); return out

NEW_INTRO =parse_tex('/home/user/MP3Translator/rapport/_intro.tex')
NEW_CH4   =parse_tex('/home/user/MP3Translator/rapport/_ch4.tex')
NEW_THEORY=parse_tex('/home/user/MP3Translator/rapport/_theory.tex')

# ---------- low-level docx helpers ----------
def _set(el,tag,**a):
    e=OxmlElement(tag)
    for k,v in a.items(): e.set(qn(k),v)
    el.append(e); return e

def style_run(run,size,bold=False):
    run.font.name=AR; run.font.size=Pt(size); run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for k in ('w:cs','w:ascii','w:hAnsi'): rf.set(qn(k),AR)
    szcs=rpr.find(qn('w:szCs'))
    if szcs is None: szcs=_set(rpr,'w:szCs')
    szcs.set(qn('w:val'),str(int(size*2)))
    if rpr.find(qn('w:rtl')) is None: _set(rpr,'w:rtl',**{'w:val':'1'})

def rtl_par(p):
    ppr=p._p.get_or_add_pPr()
    if ppr.find(qn('w:bidi')) is None: _set(ppr,'w:bidi',**{'w:val':'1'})

def fmt_par(p,align,size,indent=False,ls=1.5,sb=0,sa=6):
    rtl_par(p); pf=p.paragraph_format
    pf.line_spacing=ls; pf.space_before=Pt(sb); pf.space_after=Pt(sa)
    pf.alignment={'j':WD_ALIGN_PARAGRAPH.JUSTIFY,'c':WD_ALIGN_PARAGRAPH.CENTER,
                  'r':WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.first_line_indent=Cm(1.27) if indent else None

def add_runs(p,text,size,bold_all=False):
    for seg in re.split(r'(⟦b⟧|⟦/b⟧)',text):
        if seg=='⟦b⟧': bold_all=True; continue
        if seg=='⟦/b⟧': bold_all=False; continue
        if seg: style_run(p.add_run(seg),size,bold_all)

def set_outline(p,lvl):
    ppr=p._p.get_or_add_pPr()
    ol=ppr.find(qn('w:outlineLvl'))
    if ol is None: ol=_set(ppr,'w:outlineLvl')
    ol.set(qn('w:val'),str(lvl))

HSIZE={'h1':16,'h2':15,'h3':14}
def build_before(ref,kind,text):
    p=ref.insert_paragraph_before()
    if kind in HSIZE:
        sz=HSIZE[kind]; lvl={'h1':0,'h2':1,'h3':2}[kind]
        fmt_par(p,'c' if kind=='h1' else 'r',sz,indent=False,ls=1.0,
                sb=(16 if kind=='h1' else 8),sa=6)
        set_outline(p,lvl); add_runs(p,text,sz,bold_all=True)
        if kind=='h1': p.paragraph_format.page_break_before=True
        else: p.paragraph_format.keep_with_next=True
    elif kind=='li':
        fmt_par(p,'j',14,indent=False,sa=3); add_runs(p,'•  '+text,14)
    else:
        fmt_par(p,'j',14,indent=True); add_runs(p,text,14)
    return p

# ---------- open + locate ----------
doc=Document(SRC)
def find(pred):
    for p in doc.paragraphs:
        if pred(p.text.strip()): return p
    return None
intro_h=find(lambda t:t.startswith('المقدمة العامة'))
ch1_h  =find(lambda t:t.startswith('الفصل الاول') or t.startswith('الفصل الأول'))
ch4_h  =find(lambda t:t.startswith('الفصل الرابع'))
khat_h =find(lambda t:t.startswith('الخاتمة العامة'))
shukr_h=find(lambda t:t.startswith('شكر وتقدير'))

def delete_between(startp,endp):
    cur=startp._p; end=endp._p
    while cur is not None and cur is not end:
        nxt=cur.getnext(); cur.getparent().remove(cur); cur=nxt

# remove the original manual "لائحة الجداول" (no page numbers) — replaced by an auto one
lot_manual=find(lambda t:t.startswith('لائحة الجداول'))
khulasa   =find(lambda t:t.startswith('الخلاصة'))
if lot_manual is not None and khulasa is not None:
    delete_between(lot_manual,khulasa)

# (intro & Chapter-4 replacement happens AFTER normalisation, see below)

# ---------- cover: Faculté des Sciences, Tétouan ----------
def para_index(pred):
    for i,p in enumerate(doc.paragraphs):
        if pred(p.text.strip()): return i
    return 0
_cover_end=para_index(lambda t:t.startswith('شكر وتقدير'))
for p in doc.paragraphs[:_cover_end]:
    if 'كلية الآداب' in p.text or 'بمرتيل' in p.text or 'العلوم الإنسانية' in p.text:
        for r in p.runs:
            r.text=(r.text.replace('كلية الآداب والعلوم الإنسانية بمرتيل','كلية العلوم بتطوان')
                          .replace('كلية الآداب والعلوم الإنسانية','كلية العلوم')
                          .replace('بمرتيل','بتطوان').replace('مرتيل','تطوان'))

# ---------- margins + RTL on every section (page numbering set later) ----------
for sec in doc.sections:
    sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(3); sec.right_margin=Cm(3)
    if sec._sectPr.find(qn('w:bidi')) is None: sec._sectPr.append(OxmlElement('w:bidi'))

# ---------- normalise body formatting + headings ----------
SEC_PREF=('أولاً','أولًا','ثانياً','ثانيًا','ثالثاً','ثالثًا','رابعاً','رابعًا','خامساً','خامسًا',
          'سادساً','سادسًا','سابعاً','سابعًا','ثامناً','ثامنًا','تاسعاً','تاسعًا','عاشراً')
BED=('البعد الأول','البعد الثاني','البعد الثالث','البعد الرابع','البعد الخامس')
MEHWAR=('المحور الأول','المحور الثاني','المحور الثالث')
SUBQ=('السؤال الأول','السؤال الثاني','السؤال الافتتاحي','السؤال الختامي','سؤال افتتاحي','سؤال ختامي')
H1=('الفصل','المقدمة العامة','الخاتمة العامة','الملاحق','قائمة المراجع','شكر وتقدير',
    'الخلاصة','ملخص','Résumé','لائحة الجداول','الفهرس')
def classify(t):
    tt=t.strip()
    if not tt: return None
    if any(tt.startswith(h) for h in H1): return 'h1'
    if len(tt)<70 and tt.startswith(SEC_PREF): return 'h2'
    if len(tt)<70 and any(tt.startswith(b) for b in BED): return 'h2'
    if len(tt)<85 and any(tt.startswith(m) for m in MEHWAR): return 'h2'
    if tt.startswith('الملحق'): return 'h2'
    if tt in ('تمهيد','تمهيد الفصل','الأسئلة الفرعية','السؤال المركزي'): return 'h2'
    if len(tt)<70 and any(tt.startswith(q) for q in SUBQ): return 'h3'
    if 'التعريف الإجرائي' in tt and len(tt)<45: return 'h3'
    if tt.startswith('مفهوم الإدراك'): return 'h3'
    if re.match(r'^\d+[.\-]\s',tt) and len(tt)<45: return 'h3'
    return None

paras=doc.paragraphs
cover_end=para_index(lambda t:t.startswith('شكر وتقدير'))
for idx,p in enumerate(paras):
    txt=p.text.strip()
    if idx<cover_end:                      # cover: centre, keep, just set font
        rtl_par(p); p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            sz=r.font.size.pt if r.font.size else 16
            style_run(r,min(sz,22),bold=bool(r.font.bold))
        continue
    if not txt:
        continue
    cls=classify(txt)
    islist=(p.style.name=='List Paragraph')
    if cls:
        sz=HSIZE[cls]; lvl={'h1':0,'h2':1,'h3':2}[cls]
        fmt_par(p,'c' if cls=='h1' else 'r',sz,indent=False,ls=1.0,
                sb=(16 if cls=='h1' else 8),sa=6)
        set_outline(p,lvl)
        # major point (chapter) -> new page ; sub-headings -> never stranded at page bottom
        if cls=='h1' and any(txt.startswith(m) for m in
               ('المقدمة العامة','الفصل','الخاتمة العامة','الملاحق','قائمة المراجع')):
            p.paragraph_format.page_break_before=True
        else:
            p.paragraph_format.keep_with_next=True
        for r in p.runs: style_run(r,sz,bold=True)
    else:
        fmt_par(p,'j',14,indent=(not islist))
        p.paragraph_format.widow_control=True
        for r in p.runs: style_run(r,14,bold=bool(r.font.bold))

# ---------- make every table fit the page (RTL) ----------
for t in doc.tables:
    ncols=len(t.columns)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.allow_autofit=True
    tblPr=t._tbl.tblPr
    if tblPr.find(qn('w:bidiVisual')) is None: _set(tblPr,'w:bidiVisual')
    tw=tblPr.find(qn('w:tblW'))
    if tw is None: tw=_set(tblPr,'w:tblW')
    tw.set(qn('w:type'),'pct'); tw.set(qn('w:w'),'5000')
    lay=tblPr.find(qn('w:tblLayout'))
    if lay is None: lay=_set(tblPr,'w:tblLayout')
    lay.set(qn('w:type'),'autofit')
    fsz=9 if ncols>=6 else (10 if ncols>=4 else 12)
    for row in t.rows:
        for cell in row.cells:
            for cp in cell.paragraphs:
                rtl_par(cp); cp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.line_spacing=1.0; cp.paragraph_format.space_after=Pt(0)
                cp.paragraph_format.first_line_indent=None
                for r in cp.runs: style_run(r,fsz,bold=bool(r.font.bold))

# ---------- replace intro & Chapter 4 (after normalisation, so the new,
#            already-styled paragraphs are not clobbered) ----------
delete_between(ch4_h,khat_h)
for idx,(kind,text) in enumerate(NEW_CH4):
    if idx==0 and kind=='h1': text='الباب الرابع: '+text
    build_before(khat_h,kind,text)
delete_between(intro_h,ch1_h)
mqddima=build_before(ch1_h,'h1','المقدمة العامة')
for kind,text in NEW_INTRO:
    if kind=='h1': continue          # skip its own duplicate title
    build_before(ch1_h,kind,text)

# ---------- insert the MISSING theoretical framework + Chapter 2 heading ----------
_frag=find(lambda t:t.startswith('ؤ'))
if _frag is not None: _frag._p.getparent().remove(_frag._p)
_meth=find(lambda t:t.strip().startswith('4.') and 'هدف' in t) \
      or find(lambda t:t.strip().startswith('هدف البحث'))
if _meth is not None:
    for kind,text in NEW_THEORY: build_before(_meth,kind,text)   # نهاية الفصل الأول
    build_before(_meth,'h1','الباب الثاني: الإطار المنهجي للدراسة')

# ---------- auto فهرس المحتويات + لائحة الجداول before the introduction ----------
def field_para(ref,instr):
    fp=ref.insert_paragraph_before(); rtl_par(fp)
    fp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=fp.add_run(); _set(r._r,'w:fldChar',**{'w:fldCharType':'begin'})
    r2=fp.add_run(); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=instr; r2._r.append(it)
    r3=fp.add_run(); _set(r3._r,'w:fldChar',**{'w:fldCharType':'separate'})
    r4=fp.add_run('اضغط Ctrl+A ثم F9 لتحديث الفهرس'); style_run(r4,12)
    r5=fp.add_run(); _set(r5._r,'w:fldChar',**{'w:fldCharType':'end'})
    return fp
def heading_before(ref,text):
    hp=ref.insert_paragraph_before(); fmt_par(hp,'c',16,ls=1.0,sb=14,sa=6)
    add_runs(hp,text,16,bold_all=True); return hp
def pagebreak_before(ref):
    bp=ref.insert_paragraph_before(); r=bp.add_run(); br=OxmlElement('w:br'); br.set(qn('w:type'),'page'); r._r.append(br)
pagebreak_before(mqddima)
heading_before(mqddima,'فهرس المحتويات')
field_para(mqddima,'TOC \\o "1-3" \\h \\z \\u')
heading_before(mqddima,'لائحة الجداول')
field_para(mqddima,'TOC \\h \\z \\c "جدول"')
# (no trailing page break: المقدمة العامة has page_break_before)

# ---------- page numbering: cover (none) / front matter (أ ب ج) / body (1 2 3) ----------
import copy
_ORDER_BEFORE={qn('w:cols'),qn('w:formProt'),qn('w:vAlign'),qn('w:noEndnote'),
               qn('w:titlePg'),qn('w:textDirection'),qn('w:bidi'),qn('w:rtlGutter'),
               qn('w:docGrid'),qn('w:printerSettings'),qn('w:sectPrChange')}
def set_pgnum(sectPr, fmt, start=None):
    old=sectPr.find(qn('w:pgNumType'))
    if old is not None: sectPr.remove(old)
    el=OxmlElement('w:pgNumType'); el.set(qn('w:fmt'),fmt)
    if start is not None: el.set(qn('w:start'),str(start))
    for child in sectPr:
        if child.tag in _ORDER_BEFORE:
            child.addprevious(el); break
    else:
        sectPr.append(el)
def footer_page_field(sec):
    sec.footer.is_linked_to_previous=False
    fp=sec.footer.paragraphs[0]
    for rr in list(fp.runs): rr._r.getparent().remove(rr._r)
    rtl_par(fp); fp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run(); _set(r._r,'w:fldChar',**{'w:fldCharType':'begin'})
    r2=fp.add_run(); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=' PAGE '; r2._r.append(it)
    r3=fp.add_run(); _set(r3._r,'w:fldChar',**{'w:fldCharType':'end'})
    for rr in fp.runs: style_run(rr,12)
def footer_empty(sec):
    sec.footer.is_linked_to_previous=False
    for rr in list(sec.footer.paragraphs[0].runs): rr._r.getparent().remove(rr._r)

# capture existing sectPr refs BEFORE splitting
cover_pr=doc.sections[0]._sectPr           # الغلاف
front_pr=doc.sections[1]._sectPr           # يحكم حاليًا من شكر حتى نهاية القسم الثاني
# insert a section break just before المقدمة العامة -> front matter becomes its own section
mi=[i for i,p in enumerate(doc.paragraphs) if p.text.strip().startswith('المقدمة العامة')][0]
prev=doc.paragraphs[mi-1]
new_pr=copy.deepcopy(front_pr)
for fr in new_pr.findall(qn('w:footerReference')): new_pr.remove(fr)  # own footer
prev._p.get_or_add_pPr().append(new_pr)    # section that ENDS before المقدمة = front matter
# now assign numbering formats
set_pgnum(cover_pr,'decimal',1)            # الغلاف (بدون رقم ظاهر)
set_pgnum(new_pr,'arabicAbjad',1)          # التمهيديات: أ، ب، ج ...
set_pgnum(front_pr,'decimal',1)            # المتن يبدأ من المقدمة: 1، 2، 3 ...
# footers
secs=doc.sections                          # يُعاد حسابها: [غلاف, تمهيديات, متن1, متن2...]
footer_empty(secs[0])                       # الغلاف بلا رقم
footer_page_field(secs[1])                  # التمهيديات (أبجدي)
for s in secs[2:]:
    s.footer.is_linked_to_previous=True     # المتن يرث حقل PAGE ويعرضه بالأرقام العادية
# فاصل الصفحة (nextPage) + إزالة page_break_before من المقدمة لتفادي صفحة فارغة
_t=new_pr.find(qn('w:type'))
if _t is None:
    _t=OxmlElement('w:type'); _pg=new_pr.find(qn('w:pgSz'))
    (_pg.addprevious(_t) if _pg is not None else new_pr.insert(0,_t))
_t.set(qn('w:val'),'nextPage')
doc.paragraphs[mi].paragraph_format.page_break_before=False

# ---------- fix faculty text inside text boxes (raw w:t nodes) ----------
_repl=[('كلية الآداب والعلوم الإنسانية','كلية العلوم'),
       ('الآداب والعلوم الإنسانية','العلوم'),
       ('مرتيل','تطوان'),
       ('الفصول','الأبواب'),('فصول','أبواب'),('الفصل','الباب')]  # الأقسام الكبرى = الباب
_nfix=0
for wt in doc.element.iter(qn('w:t')):
    if wt.text:
        new=wt.text
        for a,b in _repl: new=new.replace(a,b)
        if new!=wt.text: wt.text=new; _nfix+=1
print('faculty w:t nodes fixed:',_nfix)
# paragraph-level fallback for headings where الفصل is split across runs
for p in doc.paragraphs:
    full=''.join(r.text for r in p.runs)
    if 'الفصل' in full and len(full)<80:
        full=full.replace('الفصول','الأبواب').replace('فصول','أبواب').replace('الفصل','الباب')
        if p.runs:
            p.runs[0].text=full
            for r in p.runs[1:]: r.text=''

# update TOC-type fields on open (if any exist)
try:
    st=doc.settings.element; _set(st,'w:updateFields',**{'w:val':'true'})
except Exception: pass

doc.save(OUT)
print('SAVED',OUT,'| paragraphs:',len(doc.paragraphs),'| tables:',len(doc.tables))
